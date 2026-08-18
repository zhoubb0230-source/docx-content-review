#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""审查记忆：回灌人工决策 + 命中降级（spec §11.7）。

误报抑制的长期资产：有一类问题反复出现却每次都被人工驳回——公司写作习惯、
行业惯用表述、领域内约定俗成的搭配。这些不是术语，无法进 glossary。

存储位置是**工作目录级**（docx-review/review-memory.json），跨文档共享；
放 run 目录会随清理消失，失去全部价值。

**只有两条复用路径，不得有第三条**：
  1. 精确复用：键为 (category, 归一化后的 original_text) 的哈希，命中即降级为
     report_only。归一化仅限空白与全半角统一，不做任何语义处理。
  2. 人工提炼：用户把反复出现的误报提炼为 never-flag 规则或 fallback 术语。

**严禁让系统从历史决策自动泛化规则。**「用户忽略了 3 条 B3，所以以后 B3 都降级」
这类推断会制造新的、更隐蔽的漏报。记忆只能精确匹配，泛化必须经人。

子命令
  import  从标注后的 issues.xlsx 读「人工决策」列 → 写入记忆
  apply   用记忆标注 issues-verified.jsonl，命中者降级为 report_only
  show    打印记忆条目统计
退出码：0 成功；10 xlsx 不可解析。
"""
from __future__ import annotations

import argparse
import sys
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from _common import (  # noqa: E402
    EX, atomic_write_json, atomic_write_jsonl, die, emit, normalize_width, normalize_ws,
    now_iso, read_json, read_jsonl, run_cli, sha256_text,
)
from workspace import deliver_path, guard_write_path, load_run_config, resolve_path  # noqa: E402

SCHEMA_VERSION = "1"


def memory_path(run_dir: Path, cfg: dict) -> Path:
    configured = ((cfg.get("review_memory") or {}).get("path"))
    if configured:
        return Path(configured).expanduser()
    # <输出根>/docx-review/review-memory.json —— 与 run 同级之上，跨文档共享
    return resolve_path(run_dir, "review_memory")


def key_hash(category: str, original: str) -> str:
    """归一化仅限空白与全半角统一，不做任何语义处理。"""
    return sha256_text(f"{category}{normalize_ws(normalize_width(original))}")


def load_memory(path: Path) -> dict:
    m = read_json(path, None)
    if not isinstance(m, dict) or "entries" not in m:
        return {"schema_version": SCHEMA_VERSION, "entries": []}
    return m


def cmd_import(run_dir: Path, cfg: dict, xlsx: Path | None) -> dict:
    try:
        from openpyxl import load_workbook
    except ImportError:
        die(EX.ENV, "读取 issues.xlsx 需要 openpyxl（pip install openpyxl）")

    path = xlsx or deliver_path(run_dir, "issues_xlsx")
    if not path.exists():
        die(EX.USAGE, f"未找到标注后的 issues.xlsx：{path}")
    wb = load_workbook(str(path), read_only=True, data_only=True)
    ws = wb[wb.sheetnames[0]]
    rows = list(ws.iter_rows(values_only=True))
    wb.close()
    header_idx = next((i for i, r in enumerate(rows)
                       if r and "id" in [str(c).strip() if c else "" for c in r]), None)
    if header_idx is None:
        die(EX.PARSE, f"{path} 中未找到表头行（应含 id 列）")
    header = [str(c).strip() if c is not None else "" for c in rows[header_idx]]

    def col(name):
        return header.index(name) if name in header else None

    ci = {n: col(n) for n in ("id", "类别", "原文", "人工决策", "备注")}
    if ci["人工决策"] is None:
        die(EX.PARSE, f"{path} 中缺少「人工决策」列")

    mem = load_memory(memory_path(run_dir, cfg))
    index = {e["key_hash"]: e for e in mem["entries"]}
    added = updated = 0
    for r in rows[header_idx + 1:]:
        if not r:
            continue
        decision = str(r[ci["人工决策"]] or "").strip().lower()
        if decision not in ("ignore", "reject", "驳回", "忽略"):
            continue                       # 只记录驳回；accept 不产生记忆
        cat = str(r[ci["类别"]] or "").strip() if ci["类别"] is not None else ""
        orig = str(r[ci["原文"]] or "").strip() if ci["原文"] is not None else ""
        if not cat or not orig:
            continue
        kh = key_hash(cat, orig)
        note = str(r[ci["备注"]] or "").strip() if ci["备注"] is not None else ""
        if kh in index:
            index[kh]["hit_count"] = int(index[kh].get("hit_count") or 0) + 1
            updated += 1
        else:
            index[kh] = {"key_hash": kh, "category": cat,
                         "original_norm": normalize_ws(normalize_width(orig)),
                         "decision": "ignore", "reason": note or "人工驳回",
                         "created_at": now_iso(), "hit_count": 1}
            added += 1
    mem["entries"] = list(index.values())
    mem["schema_version"] = SCHEMA_VERSION
    mem["updated_at"] = now_iso()
    p = memory_path(run_dir, cfg)
    guard_write_path(p, run_dir, run_dir.parent.parent)
    atomic_write_json(p, mem)
    return {"added": added, "updated": updated, "total": len(mem["entries"]), "path": str(p)}


def cmd_apply(run_dir: Path, cfg: dict) -> dict:
    rm = cfg.get("review_memory") or {}
    if not rm.get("enabled", True):
        return {"enabled": False, "hits": 0}
    if not rm.get("exact_match_only", True):
        die(EX.USAGE, "review_memory.exact_match_only 不得关闭：记忆只能精确匹配，"
                      "泛化必须经人（spec §11.7）")
    mem = load_memory(memory_path(run_dir, cfg))
    index = {e["key_hash"]: e for e in mem["entries"]}
    src = resolve_path(run_dir, "issues_verified")
    rows = list(read_jsonl(src))
    hits = 0
    for r in rows:
        kh = key_hash(r.get("category") or "", r.get("original_text") or "")
        e = index.get(kh)
        if not e:
            r.setdefault("memory_hit", False)
            continue
        r["memory_hit"] = True
        r["memory_reason"] = e.get("reason")
        r["action"] = "report_only"        # 命中即自动降级，不进文档
        r["suggested_text"] = ""
        hits += 1
    if rows:
        guard_write_path(src, run_dir)
        atomic_write_jsonl(src, rows)
    return {"enabled": True, "memory_entries": len(index), "issues": len(rows), "hits": hits}


def main(argv: list[str]) -> int:
    ap = argparse.ArgumentParser(prog="import_decisions.py", description="审查记忆")
    sub = ap.add_subparsers(dest="cmd", required=True)
    p = sub.add_parser("import")
    p.add_argument("--run-dir", required=True)
    p.add_argument("--xlsx")
    p.add_argument("--config")
    p = sub.add_parser("apply")
    p.add_argument("--run-dir", required=True)
    p.add_argument("--config")
    p = sub.add_parser("show")
    p.add_argument("--run-dir", required=True)
    p.add_argument("--config")
    args = ap.parse_args(argv)

    run_dir = Path(args.run_dir).resolve()
    cfg = load_run_config(run_dir, args.config)
    if args.cmd == "import":
        emit({"ok": True, **cmd_import(run_dir, cfg, Path(args.xlsx) if args.xlsx else None)})
    elif args.cmd == "apply":
        emit({"ok": True, **cmd_apply(run_dir, cfg)})
    else:
        mem = load_memory(memory_path(run_dir, cfg))
        by_cat: dict[str, int] = {}
        for e in mem["entries"]:
            by_cat[e.get("category", "?")] = by_cat.get(e.get("category", "?"), 0) + 1
        emit({"ok": True, "path": str(memory_path(run_dir, cfg)),
              "entries": len(mem["entries"]), "by_category": by_cat,
              "note": "首轮为空是正常状态，价值从第二个文档开始显现"})
    return EX.OK


if __name__ == "__main__":
    run_cli(main)
