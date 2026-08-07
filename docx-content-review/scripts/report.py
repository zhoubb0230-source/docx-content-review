#!/usr/bin/env python3
# -*- coding: utf-8 -*-
"""生成 report.md / issues.xlsx（spec §10.2、§10.3）。

可从 issues-verified.jsonl + conflicts-verified.jsonl 完全重生成，无需重跑任何
Pass——这是 §11.6 故障恢复矩阵中「report.md / issues.xlsx 缺失」的恢复路径。

priority_score 完全由脚本计算（规则强度 + 类别权重 + 证据数量 + 复核一致性
+ 章节跨度），**仅用于 issues.xlsx 的展示排序，不参与任何自动判定**（D2）。
几千条问题需要排序，但排序依据不能来自模型自评——否则它会以另一个名字变回阈值。

用法
  report.py --run-dir <run> [--config]
退出码：0 成功。
"""
from __future__ import annotations

import argparse
import sys
from collections import Counter, defaultdict
from pathlib import Path

sys.path.insert(0, str(Path(__file__).resolve().parent))

from _common import (  # noqa: E402
    conflict_admitted, EX, atomic_write_text, emit, now_iso, read_json, read_jsonl, rule_label, run_cli,
    version_header,
)
from workspace import (  # noqa: E402
    deliver_meta, deliver_path, guard_write_path, load_config, resolve_path,
)

SEV_ORDER = {"Critical": 0, "High": 1, "Medium": 2, "Low": 3}
SEV_CN = {"Critical": "严重", "High": "重要", "Medium": "中等", "Low": "提示"}
CATEGORY_WEIGHT = {"A": 1.0, "B": 0.7, "C": 0.3, "L": 1.0, "P": 0.6}
COLUMNS = ["id", "严重度", "priority_score", "类别", "规则ID", "章节路径", "页码",
           "原文", "建议", "动作", "复核结果", "记忆命中", "人工决策", "备注"]


def priority_score(rec: dict) -> float:
    """脚本计算的排序信号。不得参与任何自动判定。"""
    sev = {"Critical": 4.0, "High": 3.0, "Medium": 2.0, "Low": 1.0}.get(rec.get("severity"), 1.0)
    cat = (rec.get("category") or "?")[0]
    weight = CATEGORY_WEIGHT.get(cat, 0.5)
    evidence = 1.0 if (rec.get("evidence") or rec.get("note")) else 0.0
    sides = len(rec.get("sides") or [])
    v = rec.get("verify")
    verdict = (v.get("result") if isinstance(v, dict) else v) or rec.get("verdict")
    agreement = {"pass": 1.0, "CONFLICT": 1.0, "UNSURE": 0.3}.get(verdict, 0.0)
    span = float(rec.get("chapter_span") or 0)
    return round(sev * weight + 0.5 * evidence + 0.3 * sides + 0.8 * agreement
                 + 0.4 * min(span, 3), 3)


def load_rows(run_dir: Path, cfg: dict) -> list[dict]:
    rows = []
    for r in read_jsonl(resolve_path(run_dir, "issues_verified")):
        if (r.get("verify") or {}).get("result") == "drop":
            continue
        rows.append({
            "id": r.get("id") or f"I-{len(rows)+1:05d}",
            "severity": r.get("severity") or "Medium",
            "category": r.get("category"), "rule_id": r.get("rule_id") or r.get("category"),
            "heading_path": r.get("heading_path") or [], "page_hint": r.get("page_hint"),
            "original_text": r.get("original_text") or "",
            "suggested_text": r.get("suggested_text") or "",
            "action": r.get("action") or ("revision" if r.get("suggested_text") else "comment"),
            "verify": (r.get("verify") or {}).get("result") or "n/a",
            "memory_hit": bool(r.get("memory_hit")),
            "note": r.get("gate_note") or r.get("evidence") or "",
            "pattern_name": r.get("pattern_name") or "",
            "kind": "issue",
        })

    verdicts = {v.get("conflict_id"): v for v in read_jsonl(resolve_path(run_dir, "conflicts_verified"))}
    cdir = resolve_path(run_dir, "conflicts_candidate")
    for path in sorted(cdir.glob("conflicts-candidate.*.json")):
        for c in (read_json(path, {}) or {}).get("candidates", []):
            v = verdicts.get(c["conflict_id"])
            ok, why = conflict_admitted(c, v)
            if not ok and why == "裁定为不构成矛盾":
                continue                       # 明确判否的才丢
            # 未裁定 / 非 Critical 的 UNSURE 不进交付物，但**必须在报告里可见**——
            # 静默隐藏和 fail-open 一样糟，用户会以为文档里已经没有这些冲突了
            sides = c.get("sides") or [{}]
            rows.append({
                "id": c["conflict_id"], "severity": c.get("severity") or "Medium",
                "category": c.get("rule"), "rule_id": c.get("rule"),
                "heading_path": sides[0].get("heading_path") or [],
                "page_hint": sides[0].get("page_hint"),
                "original_text": sides[0].get("text") or "",
                "suggested_text": (c.get("suggest") or {}).get("suggested_text") or "",
                "action": c.get("action") or "comment",
                "verify": (v or {}).get("verdict") or "n/a",
                "memory_hit": False, "note": c.get("note") or c.get("description") or "",
                "kind": "conflict", "sides": sides, "chapter_span": c.get("chapter_span"),
                "description": c.get("description"),
                "admitted": ok, "withheld_reason": "" if ok else why,
            })

    for r in rows:
        r["priority_score"] = priority_score(r)
    rows.sort(key=lambda r: (SEV_ORDER.get(r["severity"], 9), -r["priority_score"], r["id"]))
    return rows


def write_xlsx(run_dir: Path, rows: list[dict]) -> str | None:
    try:
        from openpyxl import Workbook
        from openpyxl.styles import Alignment, Font
    except ImportError:
        return None
    wb = Workbook()
    ws = wb.active
    ws.title = "issues"
    vh = version_header()
    # 版本标识写入头部：下游工具据 schema_version 判断兼容性，避免静默错列
    ws.append([f"spec_version={vh['spec_version']}", f"schema_version={vh['schema_version']}",
               f"skill_version={vh['skill_version']}",
               f"runid={(read_json(resolve_path(run_dir,'manifest'),{}) or {}).get('runid','')}"])
    ws.append(COLUMNS)
    for c in ws[2]:
        c.font = Font(bold=True)
    for r in rows:
        ws.append([
            r["id"], SEV_CN.get(r["severity"], r["severity"]), r["priority_score"],
            r["category"], r["rule_id"], " > ".join(r["heading_path"]), r.get("page_hint"),
            r["original_text"][:800], r["suggested_text"][:800], r["action"],
            r["verify"], "是" if r["memory_hit"] else "", "", r["note"][:400],
        ])
    widths = [12, 8, 8, 8, 8, 32, 6, 46, 46, 12, 12, 8, 10, 40]
    for i, w in enumerate(widths, 1):
        ws.column_dimensions[ws.cell(row=2, column=i).column_letter].width = w
    for row in ws.iter_rows(min_row=3):
        for c in row:
            c.alignment = Alignment(vertical="top", wrap_text=True)
    ws.freeze_panes = "A3"
    path = deliver_path(run_dir, "issues_xlsx")
    guard_write_path(path, run_dir)
    path.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(path))
    return str(path)


def render_markdown(run_dir: Path, rows: list[dict], cfg: dict) -> str:
    man = read_json(resolve_path(run_dir, "manifest"), {}) or {}
    idx = read_json(resolve_path(run_dir, "chunk_index"), {}) or {}
    heads = read_json(resolve_path(run_dir, "headings"), {}) or {}
    gl = read_json(resolve_path(run_dir, "glossary_merged"), {}) or {}
    metrics = read_json(deliver_path(run_dir, "metrics"), {}) or {}
    cand = read_json(resolve_path(run_dir, "glossary_candidates"), {}) or {}
    valid = read_json(resolve_path(run_dir, "work") / "validation.json", {}) or {}
    plan = read_json(resolve_path(run_dir, "patchlist"), {}) or {}
    paras = list(read_jsonl(resolve_path(run_dir, "paragraphs")))
    vh = version_header()
    L = []

    def w(s=""):
        L.append(s)

    w("---")
    for k, v in vh.items():
        w(f'{k}: "{v}"')
    w(f'runid: "{man.get("runid","")}"')
    w("---")
    w()
    w(f"# 内容审查报告：{man.get('source_name','')}")
    w()

    # 1 执行摘要
    sev_count = Counter(r["severity"] for r in rows)
    stats = man.get("stats", {})
    dens = heads.get("page_density", {}) or {}
    w("## 1. 执行摘要")
    w()
    w("| 项 | 值 |")
    w("|---|---|")
    w(f"| 文档规模 | {len(paras)} 段，{heads.get('total_chars',0)} 字，"
      f"约 {max((p.get('page_hint') or 1) for p in paras) if paras else 0} 页"
      f"（{'估算' if dens.get('estimated') else '按 app.xml 校准'}） |")
    w(f"| 分片 | {len(idx.get('chunks') or [])} 片"
      f"（{'单片模式' if idx.get('single_pass') else '多片'}），"
      f"token 计量 {idx.get('tokenizer','estimate')}"
      f"{'（估算，已按保守系数放大）' if idx.get('token_estimated') else ''} |")
    w(f"| 分片完成 | {stats.get('done',0)}/{stats.get('total_chunks',0)}，"
      f"失败 {stats.get('failed',0)}，跳过 {stats.get('skipped',0)} |")
    w(f"| 问题总数 | {len(rows)}（严重 {sev_count.get('Critical',0)}，"
      f"重要 {sev_count.get('High',0)}，中等 {sev_count.get('Medium',0)}，"
      f"提示 {sev_count.get('Low',0)}） |")
    acts = Counter(r["action"] for r in rows)
    w(f"| 交付动作 | 修订 {acts.get('revision',0)}，批注 {acts.get('comment',0)}，"
      f"仅报告 {acts.get('report_only',0)} |")
    if metrics.get("timing", {}).get("total_sec"):
        w(f"| 耗时 | 约 {metrics['timing']['total_sec'] // 60} 分钟 |")
    w()
    w("**置信策略说明**：本技能不使用模型自评置信度作为任何阈值。问题需依次通过四道闸门"
      "（类型白名单 → 原文逐字校验与分类编辑距离 → 不改清单硬过滤 → 盲测 A/B 二次复核），"
      "任一不通过即淘汰或降级为批注。跨章节冲突由脚本在结构化台账上做确定性比对得出，"
      "再由 Agent 逐条裁定真伪；不确定一律按「无问题」处理，因此本报告以精确率优先，"
      "不追求召回完备。")
    w()
    demoted = plan.get("demoted_to_comment") or []
    trunc = (metrics.get("gates") or {}).get("truncated_chunks") or 0
    if demoted or trunc or plan.get("overflow"):
        w(f"**被截断/降级**：降级为批注 {len(demoted)} 条"
          f"（其中超过 max_revisions 上限 {plan.get('overflow',0)} 条），"
          f"触及单片上报上限的分片 {trunc} 个。")
        w()
    if valid and not valid.get("pass", True):
        w(f"> ⚠️ 回写校验未通过：{', '.join(c['name'] for c in valid.get('checks',[]) if not c['pass'])}。"
          f"已回滚，仅交付报告。")
        w()
    if man.get("source_sha256_end") and man["source_sha256_end"] != man.get("source_sha256"):
        w("> ⚠️ **源文档在运行期间发生了变化**，请核查是否有其他程序改动了源文件。")
        w()

    # 2 必须人工确认项
    w("## 2. 必须人工确认项（严重级逻辑冲突）")
    w()
    crit = [r for r in rows if r["severity"] == "Critical"]
    if not crit:
        w("无。")
    else:
        w("以下冲突无法由程序判断孰对孰错，必须人工确认以哪一处为准。")
        w()
        for r in sorted(crit, key=lambda x: (x["heading_path"], x["id"])):
            w(f"### {r['id']}　{rule_label(r['rule_id'])}"
              f"{'：' + r['description'] if r.get('description') else ''}")
            w()
            w(f"- 章节：{' > '.join(r['heading_path']) or '（文档开头）'}")
            for i, s in enumerate(r.get("sides") or [], 1):
                w(f"- 位置 {i}（第 {s.get('page_hint')} 页，"
                  f"{' > '.join(s.get('heading_path') or []) or '—'}）：{s.get('text','')[:200]}")
            w(f"- 说明：{r['note']}")
            w()

    # 3 分类统计
    w("## 3. 分类统计")
    w()
    w("| 类别 | 问题类型 | 规则号 | 计数 | 主要动作 |")
    w("|---|---|---|---|---|")
    by_rule = defaultdict(list)
    for r in rows:
        by_rule[r["rule_id"]].append(r)
    for rule in sorted(by_rule):
        rs = by_rule[rule]
        kind = "语病/语义" if rule[0] in "ABC" else "逻辑一致性"
        act = Counter(x["action"] for x in rs).most_common(1)[0][0]
        w(f"| {kind} | {rule_label(rule)} | {rule} | {len(rs)} | {act} |")
    w()

    # 4 按章节明细
    w("## 4. 按章节明细")
    w()
    by_chapter = defaultdict(list)
    for r in rows:
        by_chapter[tuple(r["heading_path"])].append(r)
    for chap in sorted(by_chapter, key=lambda c: (len(c), c)):
        w(f"### {' > '.join(chap) or '（文档开头）'}")
        w()
        w("| id | 严重度 | 问题类型 | 页 | 原文 | 建议/说明 |")
        w("|---|---|---|---|---|---|")
        for r in by_chapter[chap]:
            orig = r["original_text"].replace("|", "\\|")[:60]
            tail = (r["suggested_text"] or r["note"]).replace("|", "\\|")[:80]
            w(f"| {r['id']} | {SEV_CN.get(r['severity'])} | "
              f"{rule_label(r['rule_id'])}（{r['rule_id']}） | "
              f"{r.get('page_hint') or '—'} | {orig} | {tail} |")
        w()

    # 5 术语表覆盖率
    w("## 5. 术语表覆盖率")
    w()
    entries = gl.get("entries") or []
    layers = gl.get("layers") or {}
    text_all = "\n".join(p["text"] for p in paras)
    appeared = [e for e in entries if (e.get("preferred") or e.get("key")) in text_all]
    w(f"- 术语表条目：{len(entries)} 条"
      f"（权威 {layers.get('authoritative',0)}／自动抽取 {layers.get('extracted',0)}／"
      f"兜底 {layers.get('fallback',0)}）")
    if entries:
        w(f"- 其中在文档中实际出现：{len(appeared)} 条（{len(appeared)*100//max(len(entries),1)}%）")
    covered = {(e.get("preferred") or e.get("key")) for e in entries}
    uncovered = [c for c in (cand.get("candidates") or []) if c["term"] not in covered]
    if uncovered:
        w()
        w("**文档中未被术语表覆盖的高频术语**（按频次降序，前 30 个，建议补充到权威表）：")
        w()
        w("| 术语 | 频次 | 来源 |")
        w("|---|---|---|")
        for c in sorted(uncovered, key=lambda x: -x.get("freq", 0))[:30]:
            w(f"| {c['term']} | {c.get('freq',0)} | {'、'.join(c.get('sources') or [])} |")
    clusters = {e.get("cluster_id"): [] for e in entries if e.get("cluster_id")
                and not e.get("cluster_confirmed")}
    for e in entries:
        if e.get("cluster_id") in clusters:
            clusters[e["cluster_id"]].append(e.get("preferred") or e.get("key"))
    clusters = {k: v for k, v in clusters.items() if len(v) > 1}
    if clusters:
        w()
        w("### 待确认概念族")
        w()
        w("以下写法可能指同一概念，也可能是不同事物。**未经确认前不触发任何规则**；"
          "确认后可导入权威术语表，L04/L05/L26 才会对其生效。")
        w()
        for cid, forms in sorted(clusters.items()):
            w(f"- `{cid}`：{'、'.join(forms)}")
    w()

    # 6 范式符合性
    w("## 6. 范式符合性")
    w()
    pr = cfg.get("pattern_review") or {}
    if not pr.get("enabled"):
        w("未启用（`pattern_review.enabled: false`）。"
          "范式规则是场景特定的，需按 `references/patterns.md` 提供规则包后开启。")
    else:
        applied: dict[str, int] = {}
        names: dict[str, str] = {}
        pdir = resolve_path(run_dir, "patterns")
        if pdir.exists():
            for f in sorted(pdir.glob("patterns-*.json")):
                if f.name.count(".") != 1:
                    continue
                d = read_json(f, {}) or {}
                for r in d.get("rules") or []:
                    names[r["id"]] = r.get("name") or r["id"]
                for c in d.get("candidates") or []:
                    applied[c["pattern_id"]] = applied.get(c["pattern_id"], 0) + 1
        miss = Counter(r["rule_id"] for r in rows if (r.get("category") or "") == "P1")
        for rid, n in miss.items():
            names.setdefault(rid, rid)
            applied.setdefault(rid, n)
        if not applied:
            w("已启用，但本文档未命中任何范式规则的适用场景。"
              "若与预期不符，多半是 `scope` 的定位条件过严——见 `references/patterns.md` 排障表。")
        else:
            w("| 规则 | 范式 | 适用段落 | 要件缺失 |")
            w("|---|---|---|---|")
            for rid in sorted(applied):
                w(f"| `{rid}` | {names.get(rid, rid)} | {applied[rid]} | {miss.get(rid, 0)} |")
            w()
            w("「适用段落」由脚本按规则的 `scope` 确定性定位，"
              "「要件缺失」是其中经裁定确认缺少必填要件的条数。"
              "**P 类只出批注，永不生成修订**——要件缺什么内容只有作者知道。")
    w()

    # 7 本次未覆盖范围
    w("## 7. 本次未覆盖范围")
    w()
    skip = cfg.get("skip") or {}
    n_code = sum(1 for p in paras if p["is_code"])
    n_tbl = len({p["table_id"] for p in paras if p.get("table_id")})
    over = [c for c in (idx.get("chunks") or []) if c.get("oversized")]
    skipped_chunks = [c for c in (man.get("chunks") or []) if c["status"] in ("skipped", "failed")]
    w("| 范围 | 数量 | 原因 |")
    w("|---|---|---|")
    if skip.get("code_blocks", True):
        w(f"| 代码/命令/配置段落 | {n_code} | 不改清单 N9：代码块内不做语言审查 |")
    if skip.get("tables_language_check", True):
        w(f"| 表格 | {n_tbl} | 不改清单 N11：表格不做语病审查（仍参与逻辑台账） |")
    if over:
        w(f"| 超长分片 | {len(over)} | 单块超过 token 预算，已独立成片但未再切分 |")
    if skipped_chunks:
        w(f"| 未完成分片 | {len(skipped_chunks)} | "
          f"重试耗尽或失败：{', '.join(c['id'] for c in skipped_chunks[:10])} |")
    if not (cfg.get("typo_check") or {}).get("enabled"):
        w("| 错别字专项 | — | typo_check 未启用；仅由主审查顺带发现，召回率有限 |")
    elif not (cfg.get("typo_check") or {}).get("oov_detection"):
        w("| 未登录词错别字 | — | 未登录词检测需 5 万词级词表，内置词表未达该规模；"
          "错别字召回上限 = common-typos.txt 的覆盖范围 |")
    if not pr.get("enabled"):
        w("| 场景描述范式 | — | pattern_review 未启用（需先提供范式规则包） |")
    # 未准入交付物的冲突候选必须在这里露面：它们在报告里有，但文档里没有
    held = [r for r in rows if r.get("kind") == "conflict" and r.get("admitted") is False]
    if held:
        by = Counter(r.get("withheld_reason") or "?" for r in held)
        for reason, n in by.most_common():
            w(f"| 未写入文档的冲突候选 | {n} | {reason}；仅在本报告与 issues.xlsx 中列出 |")
    if not (cfg.get("argument_review") or {}).get("enabled"):
        w("| 论证链审查 | — | argument_review 默认关闭（无客观阈值，误报不可收敛） |")
    w("| 格式/排版规范 | — | 本技能不做格式审查，且不修改任何样式（D9） |")
    w("| 事实性核查 | — | 只做文档内部自洽性，不判断与外部世界是否相符 |")
    w()

    # 8 参数快照
    w("## 8. 参数快照")
    w()
    ch = cfg.get("chunking") or {}
    w("| 参数 | 值 |")
    w("|---|---|")
    w(f"| strictness | {cfg.get('strictness')} |")
    w(f"| apply_threshold | {cfg.get('apply_threshold')}（固定，不可放宽） |")
    w(f"| 分片 | max_text_tokens={ch.get('max_text_tokens')}，"
      f"max_context_tokens={ch.get('max_context_tokens')}，"
      f"max_issues_per_chunk={ch.get('max_issues_per_chunk')} |")
    src = gl.get("sources") or {}
    src_desc = "、".join(f"{k}:{v.get('entries', 0)} 条" for k, v in src.items()) or "仅自动抽取"
    w(f"| 术语表来源 | {src_desc} |")
    w(f"| 运行时间 | {man.get('created_at','')} → {now_iso()} |")
    w(f"| 源文档 sha256 | `{(man.get('source_sha256') or '')[:16]}…`（全程只读，未被修改） |")
    w(f"| 交付目录 | `{man.get('deliver_dir','')}` |")
    w(f"| 临时目录（中间件，可删） | `{man.get('temp_root','')}` |")
    w()
    w("---")
    w()
    w("> **如何让下一轮更准**：本次自建的术语表已输出为 `output/glossary.merged.json`。"
      "人工修订后（补充标准写法、别名、禁用写法），下次运行时通过 "
      "`--authoritative <文件>` 传入，误报率会显著下降，且命名不一致类问题可从"
      "「批注」升级为可直接接受的「修订」。这是从「无参照物」走向「有参照物」的唯一路径。")
    w()
    w("> 评审完成后，可在 `issues.xlsx` 的「人工决策」列填写 `accept` / `ignore`，"
      "再执行 `import_decisions.py` 回灌审查记忆；被驳回的表述在后续文档中会自动降级，"
      "不再重复打扰。首轮该文件为空是正常状态，价值从第二个文档开始显现。")
    return "\n".join(L) + "\n"


def main(argv: list[str]) -> int:
    ap = argparse.ArgumentParser(prog="report.py", description="生成审查报告")
    ap.add_argument("--run-dir", required=True)
    ap.add_argument("--config")
    args = ap.parse_args(argv)
    run_dir = Path(args.run_dir).resolve()
    cfg = load_config(args.config)

    out_cfg = cfg.get("output") or {}
    rows = load_rows(run_dir, cfg)
    md = render_markdown(run_dir, rows, cfg)

    # 交付物一律写交付目录（工作目录），不留在临时目录内
    path = deliver_path(run_dir, "report")
    guard_write_path(path, run_dir)
    path.parent.mkdir(parents=True, exist_ok=True)
    atomic_write_text(path, md)
    xlsx = write_xlsx(run_dir, rows)

    gl = read_json(resolve_path(run_dir, "glossary_merged"), {})
    gl_path = None
    if gl:
        out = deliver_path(run_dir, "glossary_out")
        guard_write_path(out, run_dir)
        atomic_write_text(out, __import__("json").dumps(gl, ensure_ascii=False, indent=2) + "\n")
        gl_path = str(out)

    from workspace import artifact_all

    emit({"ok": True, "report": str(path), "issues_xlsx": xlsx, "glossary": gl_path,
          "deliver_dir": deliver_meta(run_dir)["deliver_dir"],
          "artifacts": artifact_all(run_dir), "rows": len(rows),
          "critical": sum(1 for r in rows if r["severity"] == "Critical"),
          "xlsx_skipped": None if xlsx else "openpyxl 未安装，跳过 issues.xlsx"})
    return EX.OK


if __name__ == "__main__":
    run_cli(main)
